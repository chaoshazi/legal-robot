"""Knowledge base ingestion — embeds documents and stores in Qdrant."""

from pathlib import Path

from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_qdrant import QdrantVectorStore
from qdrant_client import QdrantClient
from qdrant_client.http.models import Distance, VectorParams

from app.core.config import get_settings
from app.rag.embeddings import get_embeddings

settings = get_settings()

_EMBEDDING_DIMS: dict[str, int] = {
    "mxbai-embed-large": 1024,
    "nomic-embed-text": 768,
    "all-minilm:l6-v2": 384,
}


def _get_embedding_dim() -> int:
    """Detect the embedding dimension for the configured model."""
    model = settings.ollama_embed_model
    if model in _EMBEDDING_DIMS:
        return _EMBEDDING_DIMS[model]
    emb = get_embeddings(model)
    return len(emb.embed_query("x"))


def _ensure_collection(client: QdrantClient):
    """Create the Qdrant collection if it doesn't exist."""
    collections = client.get_collections().collections
    if not any(c.name == settings.qdrant_collection for c in collections):
        dim = _get_embedding_dim()
        client.create_collection(
            collection_name=settings.qdrant_collection,
            vectors_config=VectorParams(size=dim, distance=Distance.COSINE),
        )


# ── Multi-format text extraction ────────────────────────────────────────

_SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx", ".doc", ".xlsx"}


def _extract_text(filepath: str | Path) -> str:
    """Extract text from a file based on its extension.

    Falls back to OCR for scanned PDFs (text extraction yields <50 chars).
    """
    path = Path(filepath)
    ext = path.suffix.lower()

    if ext == ".txt" or ext == ".md":
        return path.read_text("utf-8")

    if ext == ".pdf":
        return _extract_pdf(path)

    if ext == ".docx" or ext == ".doc":
        return _extract_docx(path)

    if ext == ".xlsx":
        return _extract_xlsx(path)

    raise ValueError(f"Unsupported file extension: {ext}")


def _extract_pdf(path: Path) -> str:
    """Extract text from a PDF. Falls back to OCR if the PDF appears scanned."""
    text = ""
    try:
        import pypdf

        reader = pypdf.PdfReader(str(path))
        for page in reader.pages:
            page_text = page.extract_text() or ""
            text += page_text + "\n"
    except Exception:
        text = ""

    # If text extraction yielded nothing meaningful, assume scanned and OCR
    if len(text.strip()) < 50:
        ocr_text = _ocr_pdf(path)
        if ocr_text.strip():
            return ocr_text

    return text.strip()


def _ocr_pdf(path: Path) -> str:
    """Convert PDF pages to images and run PaddleOCR."""
    try:
        import pdf2image
        from paddleocr import PaddleOCR

        ocr = PaddleOCR(use_angle_cls=True, lang="ch", show_log=False)
        images = pdf2image.convert_from_path(str(path), dpi=200)
    except Exception:
        return ""

    lines: list[str] = []
    for img in images:
        import numpy as np

        img_array = np.array(img)
        result = ocr.ocr(img_array, cls=True)
        if result and result[0]:
            page_text = " ".join(item[1][0] for item in result[0])
            lines.append(page_text)
    return "\n".join(lines)


def _extract_docx(path: Path) -> str:
    """Extract text from a .docx or .doc file."""
    import docx

    try:
        doc = docx.Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs)
    except Exception:
        pass

    # Fallback for old .doc format — use antiword
    try:
        import subprocess
        result = subprocess.run(
            ["/root/bin/antiword", str(path)],
            capture_output=True, text=True, timeout=30,
        )
        if result.returncode == 0:
            return result.stdout
    except Exception:
        pass

    return ""



def _extract_xlsx(path: Path) -> str:
    """Extract text from an .xlsx file — reads all cell values."""
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append(" | ".join(cells))
    wb.close()
    return "\n".join(lines)


# ── Ingestion ───────────────────────────────────────────────────────────


async def ingest_file(filepath: str | Path, doc_id: str | None = None) -> int:
    """Ingest a single file into Qdrant. Returns chunk count."""
    text = _extract_text(filepath)
    if not text.strip():
        return 0

    filename = Path(filepath).name

    splitter = RecursiveCharacterTextSplitter(chunk_size=400, chunk_overlap=50)
    chunks = splitter.split_text(text)

    client = QdrantClient(url=settings.qdrant_url)
    _ensure_collection(client)

    embeddings = get_embeddings()

    vector_store = QdrantVectorStore(
        client=client,
        collection_name=settings.qdrant_collection,
        embedding=embeddings,
    )
    metadata: dict[str, str] = {"source": filename}
    if doc_id:
        metadata["doc_id"] = doc_id
    metadatas = [dict(metadata)] * len(chunks)
    vector_store.add_texts(chunks, metadatas=metadatas)
    return len(chunks)


async def ingest_all() -> int:
    """Ingest all supported files in the data directory. Returns total chunks."""
    data_dir = Path(__file__).resolve().parent.parent.parent / "data"
    total = 0
    if not data_dir.exists():
        return 0

    for ext in _SUPPORTED_EXTENSIONS:
        for filepath in sorted(data_dir.glob(f"*{ext}")):
            total += await ingest_file(str(filepath))
    return total


if __name__ == "__main__":
    import asyncio
    total = asyncio.run(ingest_all())
    print(f"Ingested {total} chunks total.")
