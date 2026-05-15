"""Shared file processing utilities — text extraction, OCR, image analysis."""

from pathlib import Path


SUPPORTED_DOCUMENT_EXTENSIONS = {".txt", ".md", ".pdf", ".docx", ".doc", ".xlsx"}
SUPPORTED_IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".bmp"}
SUPPORTED_AUDIO_EXTENSIONS = {".webm", ".wav", ".mp3", ".ogg", ".m4a", ".mp4"}


def detect_file_type(mime_type: str, filename: str) -> str:
    """Return 'image', 'document', or 'audio' based on MIME type and extension."""
    if mime_type.startswith("image/"):
        return "image"
    if mime_type.startswith("audio/"):
        return "audio"
    ext = Path(filename).suffix.lower()
    if ext in SUPPORTED_IMAGE_EXTENSIONS:
        return "image"
    if ext in SUPPORTED_AUDIO_EXTENSIONS:
        return "audio"
    return "document"


def extract_text_from_file(filepath: str | Path) -> str:
    """Extract text from a document file based on its extension.

    Falls back to OCR for scanned PDFs (text extraction yields <50 chars).
    """
    path = Path(filepath)
    ext = path.suffix.lower()

    if ext in (".txt", ".md"):
        return path.read_text("utf-8")

    if ext == ".pdf":
        return _extract_pdf(path)

    if ext in (".docx", ".doc"):
        return _extract_docx(path)

    if ext == ".xlsx":
        return _extract_xlsx(path)

    raise ValueError(f"Unsupported file extension: {ext}")


def extract_text_from_image(filepath: str | Path) -> str:
    """Extract text from an image using Tesseract OCR."""
    path = Path(filepath)
    if not path.exists():
        return ""

    try:
        import pytesseract
        from PIL import Image

        img = Image.open(str(path))
        text = pytesseract.image_to_string(img, lang="chi_sim+eng")
        return text.strip()
    except Exception:
        import logging
        logging.getLogger("file_processing").exception("Image OCR failed")

    return ""


def _extract_pdf(path: Path) -> str:
    """Extract text from a PDF. Falls back to OCR if the PDF appears scanned."""
    text = ""
    try:
        import pypdf

        reader = pypdf.PdfReader(str(path))
        for page in reader.pages:
            page_text = page.extract_text() or ""
            text += page_text + "\n"
    except Exception:
        text = ""

    if len(text.strip()) < 50:
        ocr_text = _ocr_pdf(path)
        if ocr_text.strip():
            return ocr_text

    return text.strip()


def _ocr_pdf(path: Path) -> str:
    """Convert PDF pages to images and run Tesseract OCR."""
    try:
        import pdf2image
        import pytesseract

        images = pdf2image.convert_from_path(str(path), dpi=200)
    except Exception:
        return ""

    lines: list[str] = []
    for img in images:
        text = pytesseract.image_to_string(img, lang="chi_sim+eng")
        if text.strip():
            lines.append(text.strip())
    return "\n".join(lines)


def _extract_docx(path: Path) -> str:
    """Extract text from a .docx or .doc file."""
    import docx

    try:
        doc = docx.Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs)
    except Exception:
        pass

    try:
        import subprocess
        result = subprocess.run(
            ["/root/bin/antiword", str(path)],
            capture_output=True, text=True, timeout=30,
        )
        if result.returncode == 0:
            return result.stdout
    except Exception:
        pass

    return ""


def _extract_xlsx(path: Path) -> str:
    """Extract text from an .xlsx file — reads all cell values."""
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append(" | ".join(cells))
    wb.close()
    return "\n".join(lines)
